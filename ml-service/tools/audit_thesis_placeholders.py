#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from datetime import datetime

from docx import Document


EXPECTED_TODOS = {
    "TODO_AD_SOYAD",
    "TODO_DANISMAN",
    "TODO_JURI_1",
    "TODO_JURI_2",
    "TODO_JURI_3",
    "TODO_JURI_4",
    "TODO_JURI_5",
}

PATTERNS = [
    r"TODO_[A-Z0-9_]+",
    r"\bTODO\b",
    r"\bFIXME\b",
    r"\bXXX\b",
    r"placeholder",
    r"doldurulacak",
    r"düzenlenecektir",
    r"daha sonra",
]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", default="docs/tez_ana_taslak_tr_guncel_sau_fbe.docx")
    parser.add_argument("--output-md", default="docs/thesis_placeholder_audit.md")
    parser.add_argument("--output-json", default="docs/thesis_placeholder_audit.json")
    args = parser.parse_args()

    docx_path = Path(args.docx)
    if not docx_path.exists():
        raise SystemExit(f"[ERROR] DOCX not found: {docx_path}")

    doc = Document(docx_path)

    hits = []

    for idx, para in enumerate(doc.paragraphs, start=1):
        text = para.text.strip()
        if not text:
            continue

        for pat in PATTERNS:
            if re.search(pat, text, flags=re.IGNORECASE):
                tokens = re.findall(r"TODO_[A-Z0-9_]+", text)
                expected_only = bool(tokens) and all(t in EXPECTED_TODOS for t in tokens)
                status = "expected_frontmatter_todo" if expected_only else "review"
                hits.append({
                    "paragraph": idx,
                    "pattern": pat,
                    "status": status,
                    "text": text,
                })
                break

    review_hits = [h for h in hits if h["status"] == "review"]
    expected_hits = [h for h in hits if h["status"] == "expected_frontmatter_todo"]

    summary = {
        "generated_at_utc": datetime.utcnow().isoformat(),
        "docx": str(docx_path),
        "hit_count": len(hits),
        "expected_frontmatter_todo_count": len(expected_hits),
        "review_count": len(review_hits),
        "hits": hits,
    }

    md = []
    md.append("# Thesis Placeholder Audit")
    md.append("")
    md.append(f"- Generated at UTC: `{summary['generated_at_utc']}`")
    md.append(f"- DOCX: `{docx_path}`")
    md.append(f"- Hit count: `{len(hits)}`")
    md.append(f"- Expected frontmatter TODO count: `{len(expected_hits)}`")
    md.append(f"- Review count: `{len(review_hits)}`")
    md.append("")

    md.append("## 1. Review Required")
    md.append("")
    if review_hits:
        md.append("| Paragraph | Pattern | Text |")
        md.append("|---:|---|---|")
        for h in review_hits:
            text = h["text"].replace("|", "\\|")
            md.append(f"| {h['paragraph']} | `{h['pattern']}` | {text[:220]} |")
    else:
        md.append("No unexpected placeholders detected.")

    md.append("")
    md.append("## 2. Expected Frontmatter TODOs")
    md.append("")
    if expected_hits:
        md.append("| Paragraph | Text |")
        md.append("|---:|---|")
        for h in expected_hits:
            text = h["text"].replace("|", "\\|")
            md.append(f"| {h['paragraph']} | {text[:220]} |")
    else:
        md.append("No expected frontmatter TODOs detected.")

    Path(args.output_md).write_text("\n".join(md), encoding="utf-8")
    Path(args.output_json).write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print("[INFO] MD:", args.output_md)
    print("[INFO] JSON:", args.output_json)
    print("[INFO] Hit count:", len(hits))
    print("[INFO] Expected frontmatter TODO count:", len(expected_hits))
    print("[INFO] Review count:", len(review_hits))


if __name__ == "__main__":
    main()
